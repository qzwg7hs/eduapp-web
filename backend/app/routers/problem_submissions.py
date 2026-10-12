# -*- coding: utf-8 -*-
"""
Students can author their own problem (text/formula/image/heading blocks,
same shape as Lesson.content_blocks) and send it in as a suggestion. Purely
a review inbox for admin — nothing here ever auto-publishes into real
content; see models.ProblemSubmission for the full reasoning.
"""
import io
import urllib.request
from uuid import UUID

from docx import Document
from docx.shared import Inches, Pt
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Profile, ProblemSubmission
from ..schemas import (
    ProblemSubmissionAdminOut, ProblemSubmissionCreate, ProblemSubmissionExportRequest, ProblemSubmissionOut,
)
from ..auth import require_admin, require_student
from ..storage import upload_file_bytes
from datetime import datetime

router = APIRouter(prefix="/problem-submissions", tags=["problem-submissions"])

_IMG_EXT = {
    'image/png': '.png', 'image/jpeg': '.jpg', 'image/jpg': '.jpg',
    'image/gif': '.gif', 'image/webp': '.webp',
}


def require_submission_access(current_user: Profile = Depends(require_student)) -> Profile:
    """Same as require_student, plus the (rare, admin-set) per-student
    opt-out — mirrors require_duel_access in routers/duels.py."""
    if current_user.submission_disabled:
        raise HTTPException(403, "submission_access_disabled")
    return current_user


# ── Student ─────────────────────────────────────────────────────────────────

@router.post("", response_model=ProblemSubmissionOut)
def create_submission(body: ProblemSubmissionCreate, db: Session = Depends(get_db), current_user: Profile = Depends(require_submission_access)):
    if not body.content_blocks:
        raise HTTPException(400, "submission_empty")
    row = ProblemSubmission(
        student_id=current_user.id,
        title=(body.title or "").strip() or None,
        content_blocks=body.content_blocks,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


@router.get("/mine", response_model=list[ProblemSubmissionOut])
def list_mine(db: Session = Depends(get_db), current_user: Profile = Depends(require_submission_access)):
    return (
        db.query(ProblemSubmission)
        .filter(ProblemSubmission.student_id == current_user.id)
        .order_by(ProblemSubmission.created_at.desc())
        .all()
    )


@router.post("/upload-image")
async def upload_submission_image(file: UploadFile = File(...), _=Depends(require_submission_access)):
    ctype = file.content_type or ''
    if not ctype.startswith('image/'):
        raise HTTPException(400, "Only image files are supported")
    content = await file.read()
    ext = _IMG_EXT.get(ctype, '.png')
    url = upload_file_bytes(content, f'submission{ext}', folder='submission-images/', content_type=ctype)
    return {"url": url}


# ── Admin ───────────────────────────────────────────────────────────────────

def _to_admin_out(row: ProblemSubmission, student: Profile) -> ProblemSubmissionAdminOut:
    return ProblemSubmissionAdminOut(
        id=row.id, student_id=row.student_id,
        student_name=student.name, student_surname=student.surname, student_unique_id=student.unique_id,
        title=row.title, content_blocks=row.content_blocks, created_at=row.created_at, is_seen=row.is_seen,
    )


@router.get("/admin", response_model=list[ProblemSubmissionAdminOut])
def admin_list(db: Session = Depends(get_db), _=Depends(require_admin)):
    rows = (
        db.query(ProblemSubmission, Profile)
        .join(Profile, Profile.id == ProblemSubmission.student_id)
        .order_by(ProblemSubmission.created_at.desc())
        .all()
    )
    return [_to_admin_out(sub, student) for sub, student in rows]


@router.get("/admin/{submission_id}", response_model=ProblemSubmissionAdminOut)
def admin_get(submission_id: UUID, db: Session = Depends(get_db), _=Depends(require_admin)):
    row = db.query(ProblemSubmission).filter(ProblemSubmission.id == submission_id).first()
    if not row:
        raise HTTPException(404, "Not found")
    if not row.is_seen:
        row.is_seen = True
        row.seen_at = datetime.utcnow()
        db.commit()
        db.refresh(row)
    student = db.query(Profile).filter(Profile.id == row.student_id).first()
    return _to_admin_out(row, student)


@router.delete("/admin/{submission_id}")
def admin_delete(submission_id: UUID, db: Session = Depends(get_db), _=Depends(require_admin)):
    row = db.query(ProblemSubmission).filter(ProblemSubmission.id == submission_id).first()
    if not row:
        raise HTTPException(404, "Not found")
    db.delete(row)
    db.commit()
    return {"ok": True}


def _fetch_image_bytes(url: str) -> bytes | None:
    # The R2 public endpoint 403s urllib's default "Python-urllib/x.y"
    # User-Agent (bot-blocking) — a normal browser-like one goes through fine.
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            return resp.read()
    except Exception:
        return None  # a broken/unreachable image shouldn't fail the whole export


def _build_export_docx(rows: list[tuple[ProblemSubmission, Profile]]) -> bytes:
    """One numbered entry per submission ("1. Name Surname", then its
    content in order) — matches the format asked for. Formulas come out as
    plain LaTeX text, not a native editable Word equation: python-docx has
    no OMML/equation support, and converting arbitrary LaTeX to Word's math
    XML correctly would need a real LaTeX->OMML parser, which is its own
    project — explicitly allowed to fall back to LaTeX text instead."""
    doc = Document()
    total = len(rows)
    for idx, (sub, student) in enumerate(rows, start=1):
        heading = doc.add_paragraph()
        run = heading.add_run(f"{idx}. {student.name} {student.surname}")
        run.bold = True
        run.font.size = Pt(13)

        if sub.title:
            tp = doc.add_paragraph()
            trun = tp.add_run(sub.title)
            trun.italic = True

        for block in (sub.content_blocks or []):
            btype = block.get('type') if isinstance(block, dict) else None
            if btype == 'heading':
                level = block.get('level') or 2
                doc.add_heading(block.get('content') or '', level=min(max(int(level), 1), 3))
            elif btype == 'text':
                doc.add_paragraph(block.get('content') or '')
            elif btype == 'formula':
                fp = doc.add_paragraph()
                frun = fp.add_run(f"LaTeX: {block.get('latex') or ''}")
                frun.italic = True
                frun.font.name = 'Consolas'
            elif btype == 'image':
                url = block.get('url')
                img_bytes = _fetch_image_bytes(url) if url else None
                if img_bytes:
                    try:
                        doc.add_picture(io.BytesIO(img_bytes), width=Inches(4))
                    except Exception:
                        doc.add_paragraph(f"[{url}]")
                elif url:
                    doc.add_paragraph(f"[{url}]")
                caption = block.get('caption')
                if caption:
                    cp = doc.add_paragraph()
                    crun = cp.add_run(caption)
                    crun.italic = True
                    crun.font.size = Pt(9)

        if idx < total:
            doc.add_page_break()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.post("/admin/export")
def admin_export(body: ProblemSubmissionExportRequest, db: Session = Depends(get_db), _=Depends(require_admin)):
    if not body.ids:
        raise HTTPException(400, "No submissions selected")
    rows = (
        db.query(ProblemSubmission, Profile)
        .join(Profile, Profile.id == ProblemSubmission.student_id)
        .filter(ProblemSubmission.id.in_(body.ids))
        .order_by(ProblemSubmission.created_at.asc())
        .all()
    )
    content = _build_export_docx(rows)
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=problem_submissions.docx"},
    )
